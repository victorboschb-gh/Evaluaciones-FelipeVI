import json
import os
import io
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime


def generate_word_bytes(schedule_data):
    grouped_by_date = {}
    for item in schedule_data:
        date = item['date']
        if date not in grouped_by_date:
            grouped_by_date[date] = {}
        room = item['room']
        if room not in grouped_by_date[date]:
            grouped_by_date[date][room] = []
        grouped_by_date[date][room].append(item)

    doc = Document()

    title = doc.add_heading('Calendario de Evaluaciones', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for date, rooms in sorted(grouped_by_date.items()):
        date_obj = datetime.strptime(date, '%Y-%m-%d')
        date_str = date_obj.strftime('%d/%m/%Y')

        doc.add_heading(f'Fecha: {date_str}', level=1)

        for room, items in sorted(rooms.items()):
            doc.add_heading(f'Lugar: {room}', level=2)

            table = doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Hora'
            hdr_cells[1].text = 'Grupo'
            hdr_cells[2].text = 'Familia'
            hdr_cells[3].text = 'Profesores'

            items.sort(key=lambda x: x['start_time'])

            for item in items:
                row_cells = table.add_row().cells
                row_cells[0].text = f"{item['start_time']} - {item['end_time']}"
                row_cells[1].text = item['group_name']
                row_cells[2].text = item['family']
                row_cells[3].text = ', '.join(item['teachers'])

            doc.add_paragraph()

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def generate_word(json_path, output_path):
    if not os.path.exists(json_path):
        raise FileNotFoundError(f"Error: No se encontró el archivo {json_path}")

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc_bytes = generate_word_bytes(data)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, 'wb') as f:
        f.write(doc_bytes.read())

    print(f"Documento Word guardado en {output_path}")


if __name__ == '__main__':
    generate_word('.tmp/proposed_schedule.json', 'deliverables/Calendario_Final.docx')
